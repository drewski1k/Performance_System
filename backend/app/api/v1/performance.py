import uuid
from decimal import Decimal
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Company, ScorecardTemplate, ScoringPeriod, PfpConfig
from app.services.import_service import (
    detect_data_type,
    execute_hc_import,
    execute_performance_import,
    execute_unified_import,
    parse_paste,
    parse_upload,
    process_combined_data,
    process_glance_report,
    process_hc_data,
    process_qa_data,
    process_unified_data,
    validate_combined_import,
    validate_hc_import,
    validate_unified_import,
)
from app.services.scoring import calculate_scores

from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/performance", tags=["performance"])


@router.get("/template/download")
def download_import_template():
    """Download an Excel template for unified data import."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()

    # --- Instructions sheet ---
    ws_inst = wb.active
    ws_inst.title = "Instructions"
    ws_inst.column_dimensions["A"].width = 80

    instructions = [
        "PERFORMANCE DATA IMPORT TEMPLATE",
        "",
        "HOW TO USE THIS TEMPLATE:",
        "",
        "1. Go to the 'Data' sheet",
        "2. Fill in one row per agent with their roster info and metric values",
        "3. The first 5 columns (A-E) are required roster fields:",
        "   - Agent Name: Full name of the agent",
        "   - Email: Agent's email address (used to match existing agents across uploads)",
        "   - BPO: BPO partner name (e.g., Inktel, Intouch)",
        "   - Site: Site/location name",
        "   - Supervisor: Supervisor's full name",
        "",
        "4. Columns F onward are for metrics — add as many as you need:",
        "   - Each column header becomes the metric name",
        "   - Use consistent column names across uploads",
        "   - New metrics are auto-detected and added to the system",
        "   - You can configure new metrics (channel, weight, direction) in the Scorecard Config tab after upload",
        "",
        "TIPS:",
        "- Email is the primary key for matching agents across uploads",
        "- If an agent already exists (by email), their roster info will be updated",
        "- New agents are created automatically",
        "- Metric values should be numbers (no text, no symbols)",
        "- Percentage metrics: use decimals (0.92) or whole numbers (92) — be consistent",
        "",
        "EXAMPLE METRICS YOU MIGHT ADD:",
        "Voice AHT | Chat AHT | Email AHT | Voice CPH | Chat CPH",
        "QA Score | Occupancy | Productivity | Logged Hours",
        "(or any custom metrics your team tracks)",
    ]

    header_font = Font(bold=True, size=14)
    subheader_font = Font(bold=True, size=11)
    normal_font = Font(size=11)

    for i, line in enumerate(instructions, 1):
        cell = ws_inst.cell(row=i, column=1, value=line)
        if i == 1:
            cell.font = header_font
        elif line.endswith(":") or line.startswith("TIPS") or line.startswith("EXAMPLE"):
            cell.font = subheader_font
        else:
            cell.font = normal_font

    # --- Data sheet ---
    ws_data = wb.create_sheet("Data")

    headers = ["Agent Name", "Email", "BPO", "Site", "Supervisor"]
    example_metrics = ["Metric 1", "Metric 2", "Metric 3"]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    roster_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    metric_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    border = Border(
        bottom=Side(style="thin", color="B4C6E7"),
    )

    for i, h in enumerate(headers + example_metrics, 1):
        cell = ws_data.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill if i <= 5 else PatternFill(start_color="548235", end_color="548235", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        ws_data.column_dimensions[cell.column_letter].width = 18

    # Example rows
    examples = [
        ["John Smith", "john.smith@company.com", "Inktel", "Miami", "Jane Doe", 245, 180, 0.92],
        ["Maria Garcia", "maria.garcia@company.com", "Intouch", "Dallas", "Bob Wilson", 310, 95, 0.88],
        ["", "", "", "", "", "", "", ""],
    ]
    for row_idx, example in enumerate(examples, 2):
        for col_idx, val in enumerate(example, 1):
            cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
            if col_idx <= 5:
                cell.fill = roster_fill
            else:
                cell.fill = metric_fill
            cell.border = border

    # Add note row
    note_cell = ws_data.cell(row=5, column=1, value="← Add your agents below. Replace example metrics (F1, G1, H1...) with your actual metric names.")
    note_cell.font = Font(italic=True, color="666666", size=10)

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=performance_import_template.xlsx"},
    )


# --- Scoring Periods ---

class PeriodCreate(BaseModel):
    company_id: uuid.UUID
    label: str
    period_type: str
    start_date: str
    end_date: str


@router.get("/periods")
def list_periods(company_id: uuid.UUID | None = None, db: Session = Depends(get_db)):
    stmt = select(ScoringPeriod).order_by(ScoringPeriod.start_date.desc())
    if company_id:
        stmt = stmt.where(ScoringPeriod.company_id == company_id)
    return db.scalars(stmt).all()


@router.post("/periods", status_code=201)
def create_period(data: PeriodCreate, db: Session = Depends(get_db)):
    from datetime import date as date_type
    period = ScoringPeriod(
        company_id=data.company_id,
        label=data.label,
        period_type=data.period_type,
        start_date=date_type.fromisoformat(data.start_date),
        end_date=date_type.fromisoformat(data.end_date),
    )
    db.add(period)
    db.commit()
    db.refresh(period)
    return period


# --- Smart Data Import ---

class PasteData(BaseModel):
    text: str
    data_type: Optional[str] = None  # auto-detect if not provided
    cycle: Optional[int] = None
    company_name: Optional[str] = "Default Company"


class ImportExecute(BaseModel):
    text: str
    data_type: Optional[str] = None
    cycle: Optional[int] = None
    template_id: Optional[uuid.UUID] = None
    period_id: Optional[uuid.UUID] = None
    company_name: Optional[str] = "Default Company"


@router.post("/import/preview")
async def preview_upload(
    file: UploadFile,
    data_type: str | None = None,
    cycle: int | None = None,
    db: Session = Depends(get_db),
):
    """Upload a file and get a preview of what will be imported."""
    content = await file.read()
    filename = file.filename or "data.csv"
    df = parse_upload(content, filename, data_type=data_type)

    detected = data_type or detect_data_type(df)
    return _preview_dataframe(df, detected, cycle)


@router.post("/import/paste/preview")
def preview_paste(data: PasteData):
    """Paste data and get a preview of what will be imported."""
    df = parse_paste(data.text)
    detected = data.data_type or detect_data_type(df)
    return _preview_dataframe(df, detected, data.cycle)


@router.post("/import/upload")
async def execute_file_import(
    file: UploadFile,
    data_type: str | None = None,
    cycle: int | None = None,
    template_id: uuid.UUID | None = None,
    period_id: uuid.UUID | None = None,
    company_name: str = "Sephora",
    db: Session = Depends(get_db),
):
    """Upload and import a file."""
    import traceback
    content = await file.read()
    filename = file.filename or "data.csv"
    try:
        df = parse_upload(content, filename, data_type=data_type)
        detected = data_type or detect_data_type(df)
        return _execute_import(db, df, detected, cycle, template_id, period_id, company_name)
    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(500, f"Import failed: {str(e)}")


@router.post("/import/paste")
def execute_paste_import(data: ImportExecute, db: Session = Depends(get_db)):
    """Paste and import data."""
    df = parse_paste(data.text)
    detected = data.data_type or detect_data_type(df)
    return _execute_import(
        db, df, detected, data.cycle,
        data.template_id, data.period_id, data.company_name,
    )


# --- Scoring ---

@router.post("/calculate/{period_id}")
def run_scoring(
    period_id: uuid.UUID,
    template_id: uuid.UUID | None = None,
    db: Session = Depends(get_db),
):
    period = db.get(ScoringPeriod, period_id)
    if not period:
        raise HTTPException(404, "Period not found")

    # Auto-find template if not specified
    if not template_id:
        template = db.scalar(
            select(ScorecardTemplate)
            .where(ScorecardTemplate.is_active == True)
            .order_by(ScorecardTemplate.updated_at.desc())
        )
        if not template:
            template = db.scalar(select(ScorecardTemplate).limit(1))
        if not template:
            raise HTTPException(404, "No scorecard template found")
        template_id = template.id
    else:
        template = db.get(ScorecardTemplate, template_id)
        if not template:
            raise HTTPException(404, "Template not found")

    count = calculate_scores(db, period_id, template_id)
    return {"agents_scored": count, "period_id": str(period_id), "template_id": str(template_id)}


@router.post("/rescore")
def rescore_latest(
    template_id: uuid.UUID | None = None,
    period_id: uuid.UUID | None = None,
    db: Session = Depends(get_db),
):
    """Re-score agents using existing data. Auto-finds period and template if not specified."""
    # Find period
    if period_id:
        period = db.get(ScoringPeriod, period_id)
    else:
        period = db.scalar(
            select(ScoringPeriod).order_by(ScoringPeriod.start_date.desc())
        )
    if not period:
        raise HTTPException(404, "No scoring period found")

    # Find template
    if template_id:
        template = db.get(ScorecardTemplate, template_id)
    else:
        template = db.scalar(
            select(ScorecardTemplate)
            .where(ScorecardTemplate.is_active == True)
            .order_by(ScorecardTemplate.updated_at.desc())
        )
        if not template:
            template = db.scalar(select(ScorecardTemplate).limit(1))
    if not template:
        raise HTTPException(404, "No scorecard template found")

    count = calculate_scores(db, period.id, template.id)
    return {
        "agents_scored": count,
        "period_id": str(period.id),
        "period_label": period.label,
        "template_id": str(template.id),
        "template_name": template.name,
    }


# --- Helpers ---

def _preview_dataframe(df, data_type: str, cycle: int | None) -> dict:
    """Generate preview for any detected data type."""
    if data_type == "unified":
        roster, metrics = process_unified_data(df)
        result = validate_unified_import(roster, metrics)
    elif data_type == "combined":
        records = process_combined_data(df, cycle=cycle)
        result = validate_combined_import(records)
    elif data_type == "glance_report":
        records = process_glance_report(df)
        result = validate_combined_import(records)
    elif data_type == "qa_data":
        records = process_qa_data(df)
        result = validate_combined_import(records)
    elif data_type == "hc_data":
        records = process_hc_data(df)
        result = validate_hc_import(records)
    else:
        return {
            "data_type": "unknown",
            "error": "Could not detect data format. Columns found: " + ", ".join(df.columns[:20]),
            "valid_rows": 0,
        }

    result["data_type"] = data_type
    result["columns_found"] = list(df.columns)
    result["total_rows"] = len(df)
    return result


def _execute_import(
    db: Session, df, data_type: str, cycle: int | None,
    template_id: uuid.UUID | None, period_id: uuid.UUID | None,
    company_name: str,
) -> dict:
    """Execute import for any detected data type."""
    if data_type == "hc_data":
        records = process_hc_data(df)
        stats = execute_hc_import(db, records, company_name)
        return {"data_type": "hc_data", "status": "success", **stats}

    # All data types below need template + period — auto-create if missing
    if not template_id or not period_id:
        template_id, period_id = _ensure_template_and_period(db, company_name, cycle)

    if data_type == "unified":
        roster, metrics = process_unified_data(df)
        stats = execute_unified_import(
            db, roster, metrics, template_id, period_id, company_name
        )
        # Auto-run scoring after import
        scored = calculate_scores(db, period_id, template_id)
        stats["agents_scored"] = scored
        return {"data_type": "unified", "status": "success", **stats}

    if data_type == "combined":
        records = process_combined_data(df, cycle=cycle)
    elif data_type == "glance_report":
        records = process_glance_report(df)
    elif data_type == "qa_data":
        records = process_qa_data(df)
    else:
        raise HTTPException(400, f"Unknown data type: {data_type}")

    stats = execute_performance_import(db, records, template_id, period_id)

    # Auto-run scoring after import
    scored = calculate_scores(db, period_id, template_id)
    stats["agents_scored"] = scored

    return {"data_type": data_type, "status": "success", **stats}


def _ensure_template_and_period(
    db: Session, company_name: str, cycle: int | None
) -> tuple[uuid.UUID, uuid.UUID]:
    """Find or create a default scorecard template and scoring period."""
    from datetime import date, timedelta

    # Find company
    company = db.scalars(select(Company).limit(1)).first()
    if not company:
        company = Company(name=company_name)
        db.add(company)
        db.flush()

    # Find or create template
    from app.models import ScorecardMetric
    template = db.scalars(
        select(ScorecardTemplate).where(ScorecardTemplate.company_id == company.id).limit(1)
    ).first()
    if template:
        # Ensure template has metrics (fix for templates created without them)
        metric_count = db.scalar(
            select(func.count()).select_from(ScorecardMetric)
            .where(ScorecardMetric.template_id == template.id)
        )
        if not metric_count:
            _seed_template_metrics(db, template.id)
            # Also ensure PFP config exists
            existing_pfp = db.scalar(
                select(PfpConfig).where(PfpConfig.template_id == template.id)
            )
            if not existing_pfp:
                pfp = PfpConfig(
                    template_id=template.id,
                    grade_a_rate=Decimal("3"),
                    grade_b_rate=Decimal("2"),
                    grade_c_rate=Decimal("0"),
                    grade_d_rate=Decimal("0"),
                    grade_f_rate=Decimal("0"),
                )
                db.add(pfp)
                db.flush()
    if not template:
        template = ScorecardTemplate(
            company_id=company.id,
            name=f"{company.name} Scorecard",
            channel_weight=Decimal("0"),
            non_channel_weight=Decimal("100"),
            outlier_method="iqr",
            iqr_multiplier=Decimal("1.5"),
        )
        db.add(template)
        db.flush()

        # Add default scorecard metrics
        _seed_template_metrics(db, template.id)

        # Add PFP config
        pfp = PfpConfig(
            template_id=template.id,
            grade_a_rate=Decimal("3"),
            grade_b_rate=Decimal("2"),
            grade_c_rate=Decimal("0"),
            grade_d_rate=Decimal("0"),
            grade_f_rate=Decimal("0"),
        )
        db.add(pfp)
        db.flush()

    # Find or create period
    label = f"Cycle {cycle}" if cycle else "Current Period"
    period = db.scalars(
        select(ScoringPeriod)
        .where(ScoringPeriod.company_id == company.id, ScoringPeriod.label == label)
        .limit(1)
    ).first()
    if not period:
        today = date.today()
        period = ScoringPeriod(
            company_id=company.id,
            label=label,
            period_type="cycle",
            start_date=today - timedelta(days=14),
            end_date=today,
        )
        db.add(period)
        db.flush()

    return template.id, period.id


def _seed_template_metrics(db: Session, template_id: uuid.UUID) -> None:
    """Add default scorecard metrics to a new template."""
    from app.models import MetricDefinition, ScorecardMetric

    # (metric_key, weight, include_in_score, show_on_scorecard, min_threshold, threshold_basis)
    metric_configs = [
        ("voice_aht", 50, True, True, 30, "voice_contacts"),
        ("voice_cph", 50, True, True, 30, "voice_contacts"),
        ("chat_aht", 50, True, True, 30, "chat_contacts"),
        ("chat_cph", 50, True, True, 30, "chat_contacts"),
        ("email_aht", 50, True, True, 10, "email_contacts"),
        ("email_cph", 50, True, True, 10, "email_contacts"),
        ("productivity_pct", 50, True, False, 40, "logged_hours"),
        ("qa_score_pct", 50, True, False, 5, "qa_evaluations"),
        # Context-only metrics (not scored, just displayed)
        ("occupancy_pct", 0, False, True, 0, ""),
        ("total_logged_time", 0, False, True, 0, ""),
        ("total_active_time", 0, False, False, 0, ""),
        ("total_evaluations", 0, False, True, 0, ""),
        ("voice_accepted", 0, False, True, 0, ""),
        ("chat_accepted", 0, False, True, 0, ""),
        ("email_accepted", 0, False, True, 0, ""),
    ]

    for i, (key, weight, scored, show, threshold, basis) in enumerate(metric_configs):
        metric_def = db.scalar(
            select(MetricDefinition).where(MetricDefinition.key == key)
        )
        if not metric_def:
            continue
        sm = ScorecardMetric(
            template_id=template_id,
            metric_id=metric_def.id,
            weight=Decimal(str(weight)),
            include_in_score=scored,
            show_on_scorecard=show,
            min_threshold=threshold,
            threshold_basis=basis,
            grade_mode="dynamic",
            sort_order=i,
        )
        db.add(sm)
    db.flush()
